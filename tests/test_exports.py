from docx import Document
from openpyxl import load_workbook
from src.export_service import ExportService

def test_export_docx(temp_context,samples,tmp_path):
    result=temp_context.run_analysis(samples[0]); path=ExportService().export_docx(samples[0],result,tmp_path)
    assert path.exists() and 'BÁO CÁO PHÂN TÍCH TÁC ĐỘNG' in Document(path).paragraphs[0].text

def test_export_xlsx(temp_context,samples,tmp_path):
    result=temp_context.run_analysis(samples[1]); path=ExportService().export_rtm_xlsx(samples[1],result,tmp_path)
    wb=load_workbook(path); assert wb['RTM']['A1'].value=='CR ID'

def test_export_json(temp_context,samples,tmp_path):
    result=temp_context.run_analysis(samples[2]); path=ExportService().export_json(samples[2],result,tmp_path)
    assert path.exists() and 'analysis' in path.read_text(encoding='utf-8')

def test_export_csv_bytes(temp_context,samples):
    result=temp_context.run_analysis(samples[3]); data=ExportService().retrieved_csv_bytes(result)
    assert b'document_id' in data
